"""Unit tests for ExportService workbook construction.

The DB read is stubbed by replacing the service's query object with one that
returns a pre-built `EventReportContext`, so no database is needed, and the
resulting workbook is parsed back with openpyxl.

Most of what is asserted here is *absence*: a column or sheet for something
the event never used must not appear at all.
"""

from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest

from app.services.event_report_context import (
    EventAuditTrail,
    EventContent,
    EventReportContext,
    EventRoster,
)
from app.services.event_results_query import EventResultsData
from app.services.export_service import ExportService


def _team(tid: int, name: str, group: int | None = None, photo_url: str = "", total: float = 0.0):
    return SimpleNamespace(
        id=tid, name=name, versus_group_id=group, photo_url=photo_url, total=total
    )


def _cp(cid: int, order: int, name: str | None = None, **kw):
    return SimpleNamespace(
        id=cid,
        order=order,
        name=name or f"Posto {order}",
        event_id=1,
        description=kw.get("description", ""),
        clue=kw.get("clue", ""),
        staff_script=kw.get("staff_script", ""),
        challenge_brief=kw.get("challenge_brief", ""),
        latitude=kw.get("latitude"),
        longitude=kw.get("longitude"),
        arrival_radius_m=kw.get("arrival_radius_m", 50),
        available_from=None,
        available_until=None,
        is_draft=kw.get("is_draft", False),
        is_placeholder=kw.get("is_placeholder", False),
        stage_id=kw.get("stage_id"),
    )


def _result(
    team_id: int,
    checkpoint_id: int,
    *,
    final_score: float | None = 6.0,
    extra_shots: int = 0,
    penalties: dict | None = None,
    penalty_counts: dict | None = None,
    bonus_counts: dict | None = None,
    notes: str = "",
    completed: bool = True,
    judgment_status: str | None = None,
    config: dict | None = None,
    team_vs_result: str | None = None,
    result_data: dict | None = None,
    rid: int = 1,
):
    data = dict(result_data or {})
    if notes:
        data["notes"] = notes
    return SimpleNamespace(
        id=rid,
        team_id=team_id,
        activity=SimpleNamespace(
            checkpoint=SimpleNamespace(id=checkpoint_id),
            config=config or {},
            name="Prova",
            activity_type="ScoreBasedActivity",
        ),
        is_completed=completed,
        final_score=final_score,
        extra_shots=extra_shots,
        penalties=penalties or {},
        penalty_counts=penalty_counts or {},
        bonuses={},
        bonus_counts=bonus_counts or {},
        result_data=data,
        media_urls=[],
        judgment_status=judgment_status,
        team_vs_result=team_vs_result,
        completed_at=None,
    )


def _context(teams, checkpoints, results, *, roster=None, content=None, audit=None, **kwargs):
    return EventReportContext(
        results=EventResultsData(teams=teams, checkpoints=checkpoints, results=results, **kwargs),
        roster=roster or EventRoster(),
        content=content or EventContent(),
        audit=audit or EventAuditTrail(),
    )


def _service(ctx: EventReportContext) -> ExportService:
    svc = ExportService.__new__(ExportService)

    class _StubQuery:
        async def load(self, _event_id):
            return ctx

    svc.db = None
    svc._query = _StubQuery()
    return svc


async def _build(teams, checkpoints, results, **kwargs):
    ctx = _context(teams, checkpoints, results, **kwargs)
    raw = await _service(ctx).build_workbook(1)
    return openpyxl.load_workbook(BytesIO(raw))


def _rows(ws):
    return list(ws.iter_rows(values_only=True))


# ---------- structure ----------


@pytest.mark.asyncio
async def test_workbook_always_opens_with_overview_results_and_overall():
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Bar")], [_result(1, 10)])
    assert wb.sheetnames[:3] == ["Overview", "Results", "Overall"]


@pytest.mark.asyncio
async def test_checkpoint_sheets_are_named_after_the_checkpoints_that_were_used():
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Bar do Zé"), _cp(11, 2, "Praça")],
        [_result(1, 10, rid=1), _result(1, 11, rid=2)],
    )
    assert "1. Bar do Zé" in wb.sheetnames
    assert "2. Praça" in wb.sheetnames


@pytest.mark.asyncio
async def test_checkpoint_nobody_reached_gets_no_sheet_and_no_column():
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Usado"), _cp(11, 2, "Nunca aconteceu")],
        [_result(1, 10)],
    )
    assert "2. Nunca aconteceu" not in wb.sheetnames
    assert _rows(wb["Overall"])[0][:2] == ("Team", "1. Usado")


# ---------- Overall ----------


@pytest.mark.asyncio
async def test_overall_reconciles_checkpoint_subtotal_against_recorded_total():
    """Hint costs never reach a checkpoint column, so the two must be shown
    side by side rather than the report quietly ranking by the wrong one."""
    hint = SimpleNamespace(
        team_id=1, checkpoint_id=10, indication_id=1, revealed_at=datetime(2026, 5, 1), cost=-10
    )
    wb = await _build(
        [_team(1, "A", total=-4.0)],
        [_cp(10, 1, "Um")],
        [_result(1, 10, final_score=6.0)],
        hint_reveals=[hint],
    )
    header = _rows(wb["Overall"])[0]
    assert header[-4:] == ("Checkpoint Subtotal", "Adjustments", "Recorded Total", "Rank")
    row = _rows(wb["Overall"])[1]
    assert row[-3] == -10  # adjustments: the hint cost
    assert row[-2] == -4.0  # recorded total from Team.total


@pytest.mark.asyncio
async def test_overall_subtotal_and_rank_are_live_formulas():
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Um")], [_result(1, 10)])
    row = _rows(wb["Overall"])[1]
    assert str(row[-4]).startswith("=SUM(")
    assert str(row[-1]).startswith("=RANK(")


@pytest.mark.asyncio
async def test_overall_marks_a_checkpoint_the_team_never_reached_as_blank():
    wb = await _build(
        [_team(1, "A"), _team(2, "B")],
        [_cp(10, 1, "Um"), _cp(11, 2, "Dois")],
        [_result(1, 10, rid=1), _result(1, 11, rid=2), _result(2, 10, rid=3)],
    )
    rows = _rows(wb["Overall"])
    # openpyxl reads an empty-string cell back as None.
    assert rows[2][1] == 6  # B at checkpoint 1
    assert rows[2][2] is None  # B never reached checkpoint 2


@pytest.mark.asyncio
async def test_versus_column_absent_when_no_team_had_an_opponent():
    wb = await _build([_team(1, "Solo")], [_cp(10, 1, "Um")], [_result(1, 10)])
    assert "Versus Pair" not in _rows(wb["Overall"])[0]


# ---------- Results backbone ----------


@pytest.mark.asyncio
async def test_results_sheet_holds_one_row_per_result_with_its_activity():
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Bar")],
        [_result(1, 10, final_score=7.0, result_data={"achieved_points": 7})],
    )
    rows = _rows(wb["Results"])
    assert rows[0][:6] == (
        "Team",
        "Checkpoint",
        "Order",
        "Activity",
        "Activity Type",
        "Final Score",
    )
    assert rows[1][:6] == ("A", "Bar", 1, "Prova", "ScoreBasedActivity", 7)
    assert "D: achieved_points" in rows[0]


@pytest.mark.asyncio
async def test_result_data_key_no_activity_wrote_gets_no_column():
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Bar")], [_result(1, 10)])
    assert not any(str(h).startswith("D: ") for h in _rows(wb["Results"])[0])


# ---------- per-checkpoint detail ----------


@pytest.mark.asyncio
async def test_checkpoint_sheet_shows_only_the_counters_that_were_recorded():
    teams = [_team(1, "A", 5), _team(2, "B", 5)]
    results = [
        _result(1, 10, extra_shots=5, penalty_counts={"vomit": 2}, notes="Pasta", rid=1),
        _result(2, 10, final_score=4.0, rid=2),
    ]
    wb = await _build(teams, [_cp(10, 1, "Um")], results)
    header = _rows(wb["1. Um"])[0]

    assert "Vómitos (-)" in header
    assert "Extra Shots" in header
    assert "Não bebeu (-)" not in header


@pytest.mark.asyncio
async def test_peddy_paper_checkpoint_sheet_has_no_drinking_columns():
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Um")], [_result(1, 10, final_score=5.0)])
    header = _rows(wb["1. Um"])[0]
    assert header == ("Team", "Activity", "Activity Type", "Completed At", "Total Checkpoint")


@pytest.mark.asyncio
async def test_custom_activity_counter_becomes_its_own_labelled_column():
    config = {"penalty_counters": [{"key": "asneira", "label": "Asneiras", "points": 3}]}
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Um")],
        [_result(1, 10, penalty_counts={"asneira": 4}, config=config)],
    )
    assert "Asneiras (-3)" in _rows(wb["1. Um"])[0]


@pytest.mark.asyncio
async def test_global_rule_column_is_labelled_with_the_rule_name():
    rule = SimpleNamespace(id=7, name="Atraso", rule_type="penalty_counter", points=5.0)
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Um")],
        [_result(1, 10, penalty_counts={"g_7": 1})],
        rules=[rule],
    )
    assert "Atraso (-5)" in _rows(wb["1. Um"])[0]


@pytest.mark.asyncio
async def test_bonus_counter_is_shown():
    rule = SimpleNamespace(id=3, name="Criatividade", rule_type="bonus_counter", points=2.0)
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Um")],
        [_result(1, 10, bonus_counts={"gb_3": 2})],
        rules=[rule],
    )
    assert "Criatividade (+2)" in _rows(wb["1. Um"])[0]


@pytest.mark.asyncio
async def test_pending_judgment_is_labelled_rather_than_scored_as_zero():
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Um")],
        [_result(1, 10, final_score=None, judgment_status="pending_judgment")],
    )
    assert "por avaliar" in _rows(wb["1. Um"])[1]


@pytest.mark.asyncio
async def test_checkpoint_sheet_ends_with_a_totals_row_of_formulas():
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Um")], [_result(1, 10)])
    total_row = _rows(wb["1. Um"])[-1]
    assert total_row[0] == "Total"
    assert any(str(c).startswith("=SUM(") for c in total_row if c)


# ---------- route content ----------


@pytest.mark.asyncio
async def test_checkpoints_sheet_carries_the_staff_only_planning_content():
    cp = _cp(10, 1, "Bar", staff_script="Explicar as regras", challenge_brief="Beber 3 shots")
    wb = await _build([_team(1, "A")], [cp], [_result(1, 10)])
    rows = _rows(wb["Checkpoints"])
    assert "Staff Script" in rows[0]
    assert "Explicar as regras" in rows[1]
    assert "Beber 3 shots" in rows[1]


@pytest.mark.asyncio
async def test_checkpoints_sheet_aggregates_via_formulas_against_results():
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Bar")], [_result(1, 10)])
    row = _rows(wb["Checkpoints"])[1]
    assert any(str(c).startswith("=COUNTIF(") for c in row if c)
    assert any("AVERAGEIF" in str(c) for c in row if c)


@pytest.mark.asyncio
async def test_activities_sheet_absent_when_no_activities_were_configured():
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Um")], [_result(1, 10)])
    assert "Activities" not in wb.sheetnames


@pytest.mark.asyncio
async def test_activities_sheet_lists_configured_awards():
    activity = SimpleNamespace(
        id=1,
        name="Prova",
        activity_type="ScoreBasedActivity",
        checkpoint_id=10,
        config={"max_points": 100, "base_score": 50},
        is_global=False,
        is_active=True,
        available_from=None,
        available_until=None,
        description="",
    )
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Bar")],
        [_result(1, 10)],
        content=EventContent(activities=[activity]),
    )
    rows = _rows(wb["Activities"])
    assert rows[1][0] == "Bar"
    assert 100 in rows[1]


@pytest.mark.asyncio
async def test_hint_ladder_sheet_includes_the_answer_key_and_sales_count():
    indication = SimpleNamespace(
        id=1,
        checkpoint_id=10,
        order=0,
        hint="Perto do rio",
        question="Que rio?",
        expected_answer="Vouga",
    )
    reveal = SimpleNamespace(
        team_id=1, checkpoint_id=10, indication_id=1, revealed_at=datetime(2026, 5, 1), cost=-10
    )
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Bar")],
        [_result(1, 10)],
        content=EventContent(indications=[indication]),
        hint_reveals=[reveal],
    )
    row = _rows(wb["Hint Ladder"])[1]
    assert row == ("Bar", 1, 1, -10, "Perto do rio", "Que rio?", "Vouga")


@pytest.mark.asyncio
async def test_route_stage_with_no_required_count_reports_every_post():
    """NULL means all posts are required; 0 is a real value meaning none are."""
    stage = SimpleNamespace(id=1, name="Etapa A", order=1, order_matters=True, required_count=None)
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Bar", stage_id=1), _cp(11, 2, "Praça", stage_id=1)],
        [_result(1, 10), _result(1, 11, rid=2)],
        content=EventContent(stages=[stage]),
    )
    assert _rows(wb["Route Stages"])[1][3] == 2


# ---------- people ----------


@pytest.mark.asyncio
async def test_people_sheets_absent_when_nobody_was_assigned():
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Um")], [_result(1, 10)])
    for name in ("Staff", "Guides", "Team Members"):
        assert name not in wb.sheetnames


@pytest.mark.asyncio
async def test_staff_sheet_identifies_people_by_id_name_and_email():
    user = SimpleNamespace(
        id=42,
        name="Ana",
        email="ana@example.org",
        scopes=["staff"],
        disabled=False,
        authentik_sub="sub-1",
        team_id=None,
        is_captain=False,
    )
    assignment = SimpleNamespace(id=1, user_id=42, checkpoint_id=10)
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Bar")],
        [_result(1, 10)],
        roster=EventRoster(staff_assignments=[assignment], staff_users=[user]),
    )
    assert _rows(wb["Staff"])[1][:4] == (42, "Ana", "ana@example.org", "Bar")


@pytest.mark.asyncio
async def test_member_without_an_oidc_subject_is_reported_as_unlinked():
    placeholder = SimpleNamespace(
        id=7,
        name="Convidado",
        email=None,
        team_id=1,
        is_captain=False,
        authentik_sub=None,
        disabled=False,
        scopes=None,
    )
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Um")],
        [_result(1, 10)],
        roster=EventRoster(members=[placeholder]),
    )
    row = _rows(wb["Team Members"])[1]
    assert row[2] == "Convidado"
    assert row[5] == "no"  # Linked Account


@pytest.mark.asyncio
async def test_team_progress_uses_arrivals_not_the_positional_time_array():
    """`Team.times` is appended in visit order while `score_per_checkpoint` is
    rebuilt in route order, so progress must come from the arrival rows."""
    arrival = SimpleNamespace(team_id=1, checkpoint_id=11, arrived_at=datetime(2026, 5, 1, 10, 0))
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Um"), _cp(11, 2, "Dois")],
        [_result(1, 10)],
        arrivals=[arrival],
    )
    rows = _rows(wb["Team Progress"])
    assert [r[3] for r in rows[1:]] == ["Um", "Dois"]
    assert rows[2][4] == datetime(2026, 5, 1, 10, 0)


# ---------- trails ----------


@pytest.mark.asyncio
async def test_evaluation_log_absent_when_no_score_was_ever_edited():
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Um")], [_result(1, 10)])
    assert "Evaluation Log" not in wb.sheetnames


@pytest.mark.asyncio
async def test_evaluation_log_writes_one_row_per_changed_field():
    result = _result(1, 10, rid=55)
    history = SimpleNamespace(
        id=1,
        result_id=55,
        action="updated",
        editor_id="sub-1",
        editor_name="Ana",
        changes={
            "final_score": {"before": 6, "after": 9},
            "extra_shots": {"before": 0, "after": 2},
        },
        note=None,
        created_at=datetime(2026, 5, 1, 12, 0),
    )
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Bar")],
        [result],
        audit=EventAuditTrail(evaluations=[history], result_by_id={55: result}),
    )
    rows = _rows(wb["Evaluation Log"])
    assert len(rows) == 3  # header + one row per changed field
    assert {r[7] for r in rows[1:]} == {"final_score", "extra_shots"}
    assert rows[1][4] == "A"


@pytest.mark.asyncio
async def test_contest_with_no_diff_still_gets_a_row_carrying_its_reason():
    result = _result(1, 10, rid=55)
    contest = SimpleNamespace(
        id=2,
        result_id=55,
        action="contested",
        editor_id="1",
        editor_name="A",
        changes={},
        note="Não concordamos",
        created_at=datetime(2026, 5, 1, 13, 0),
    )
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Bar")],
        [result],
        audit=EventAuditTrail(evaluations=[contest], result_by_id={55: result}),
    )
    rows = _rows(wb["Evaluation Log"])
    assert len(rows) == 2
    assert rows[1][-1] == "Não concordamos"


@pytest.mark.asyncio
async def test_audit_log_sheet_appears_when_administrative_actions_were_recorded():
    entry = SimpleNamespace(
        id=1,
        event_id=1,
        actor_id="sub-1",
        actor_name="Ana",
        actor_kind="staff",
        action="rally_settings.updated",
        target_type="rally_settings",
        target_id="1",
        changes={"hints_enabled": {"before": True, "after": False}},
        note=None,
        request_id="req-1",
        created_at=datetime(2026, 5, 1, 9, 0),
    )
    wb = await _build(
        [_team(1, "A")],
        [_cp(10, 1, "Um")],
        [_result(1, 10)],
        audit=EventAuditTrail(audit_log=[entry]),
    )
    row = _rows(wb["Audit Log"])[1]
    assert row[4] == "rally_settings.updated"
    assert row[7] == "hints_enabled"
    assert (row[8], row[9]) == ("True", "False")


@pytest.mark.asyncio
async def test_side_mechanic_sheets_are_absent_when_unused():
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Um")], [_result(1, 10)])
    for name in ("Hints", "Skips", "Manual Awards", "Badges"):
        assert name not in wb.sheetnames


@pytest.mark.asyncio
async def test_manual_award_sheet_excludes_automatic_overflow_carries():
    when = datetime(2026, 5, 1, 12, 0)
    manual = SimpleNamespace(
        team_id=1,
        points=15.0,
        reason="Criatividade",
        awarded_at=when,
        activity_result_id=None,
        is_active=True,
    )
    automatic = SimpleNamespace(
        team_id=1,
        points=-4.0,
        reason="Excesso",
        awarded_at=when,
        activity_result_id=99,
        is_active=True,
    )
    wb = await _build(
        [_team(1, "A")], [_cp(10, 1, "Um")], [_result(1, 10)], awards=[manual, automatic]
    )
    rows = _rows(wb["Manual Awards"])
    assert len(rows) == 2
    assert rows[1][:3] == ("A", 15.0, "Criatividade")


# ---------- robustness ----------


@pytest.mark.asyncio
async def test_result_for_unknown_team_is_skipped():
    """A result whose team_id isn't in the current team list (e.g. the team
    was deleted after scoring) must be skipped rather than raising a
    KeyError or polluting another team's totals."""
    results = [_result(1, 10, rid=1), _result(999, 10, final_score=100.0, rid=2)]
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Um")], results)
    assert len(_rows(wb["Overall"])) == 3  # header + team A + totals row


@pytest.mark.asyncio
async def test_malformed_penalty_value_defaults_to_zero():
    results = [_result(1, 10, penalty_counts={"vomit": 2}, penalties={"vomit": "oops"})]
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Um")], results)
    assert "Vómitos (-)" in _rows(wb["1. Um"])[0]


@pytest.mark.asyncio
async def test_illegal_and_overlong_checkpoint_names_are_sanitised():
    long_name = "Bar com um nome absurdamente comprido [teste]"
    wb = await _build([_team(1, "A")], [_cp(10, 1, long_name)], [_result(1, 10)])
    title = next(t for t in wb.sheetnames if t.startswith("1."))
    assert len(title) <= 31
    assert "[" not in title
    assert "]" not in title


@pytest.mark.asyncio
async def test_skipped_checkpoint_counts_as_engagement_not_as_absence():
    """Giving up on a post is a decision the team made there, not a post it
    never reached — a blank would say the opposite of what happened."""
    skip = SimpleNamespace(
        team_id=1, checkpoint_id=10, skipped_at=datetime(2026, 5, 1, 10, 0), cost=-25
    )
    wb = await _build([_team(1, "A")], [_cp(10, 1, "Um")], [], skips=[skip])
    assert _rows(wb["Overall"])[1][1] == 0


@pytest.mark.asyncio
async def test_empty_event_still_produces_a_workbook():
    wb = await _build([], [], [])
    assert wb.sheetnames[:3] == ["Overview", "Results", "Overall"]
