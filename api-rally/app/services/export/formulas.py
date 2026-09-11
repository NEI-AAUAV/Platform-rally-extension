"""Formula strings for the workbook's totals and aggregates.

The written numbers stay exactly as the server priced them; only the sums,
averages and counts that a reader would otherwise have to redo by hand become
live formulas, so correcting a cell re-totals the sheet.

Every cross-sheet aggregate points at the flat **Results** sheet, whose title
is a constant. Checkpoint sheet titles are sanitised and truncated, so a
formula must never rebuild one by hand — pass the title `sheet_title` actually
returned, and quote it through `ref`.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter

#: The flat one-row-per-result sheet every cross-sheet aggregate reads.
RESULTS_SHEET = "Results"


def ref(sheet: str, first_row: int, last_row: int, column: int) -> str:
    """An absolute range on another sheet, with the title safely quoted.

    Excel quotes a sheet name by wrapping it in apostrophes and doubling any
    apostrophe inside it. Titles here can hold spaces and accents, so this is
    not optional.
    """
    quoted = "'" + sheet.replace("'", "''") + "'"
    letter = get_column_letter(column)
    return f"{quoted}!${letter}${first_row}:${letter}${last_row}"


def row_sum(row: int, first_col: int, last_col: int) -> str:
    """`=SUM(C2:H2)` — across one row of a sheet."""
    if last_col < first_col:
        return "=0"
    start = f"{get_column_letter(first_col)}{row}"
    end = f"{get_column_letter(last_col)}{row}"
    return f"=SUM({start}:{end})"


def col_sum(column: int, first_row: int, last_row: int) -> str:
    """`=SUM(C2:C40)` — down one column, for a totals row."""
    if last_row < first_row:
        return "=0"
    letter = get_column_letter(column)
    return f"=SUM({letter}{first_row}:{letter}{last_row})"


def col_average(column: int, first_row: int, last_row: int) -> str:
    """Average down a column, guarded so an empty range shows blank not #DIV/0!."""
    if last_row < first_row:
        return ""
    letter = get_column_letter(column)
    rng = f"{letter}{first_row}:{letter}{last_row}"
    return f'=IFERROR(AVERAGE({rng}),"")'


def rank(cell: str, column: int, first_row: int, last_row: int) -> str:
    """Descending rank of one cell within its column."""
    letter = get_column_letter(column)
    return f"=RANK({cell},${letter}${first_row}:${letter}${last_row},0)"


def count_if(criteria_range: str, criterion: str) -> str:
    return f"=COUNTIF({criteria_range},{criterion})"


def sum_if(criteria_range: str, criterion: str, sum_range: str) -> str:
    return f"=SUMIF({criteria_range},{criterion},{sum_range})"


def average_if(criteria_range: str, criterion: str, average_range: str) -> str:
    """Guarded so a row with no matching results shows blank, not an error."""
    return f'=IFERROR(AVERAGEIF({criteria_range},{criterion},{average_range}),"")'


def quote(value: str) -> str:
    """A string literal usable as a formula criterion."""
    return '"' + str(value).replace('"', '""') + '"'
