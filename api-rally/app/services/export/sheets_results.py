"""Score sheets: the event overview, the flat result list, and the breakdowns.

The **Results** sheet is the workbook's backbone: one row per persisted
`ActivityResult`, which every cross-sheet aggregate elsewhere reads through
`SUMIF`/`AVERAGEIF`. Keeping the aggregates pointed at one fixed-name sheet is
what lets the rest of the workbook use formulas without ever referencing a
checkpoint sheet whose title was sanitised and truncated.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.models.activity import ActivityResult
from app.services.event_report_columns import (
    ABSENT_XLSX,
    ReportColumn,
    checkpoint_columns,
    checkpoint_row,
)
from app.services.event_report_context import EventReportContext
from app.services.event_results_query import result_notes
from app.services.export import formulas as f
from app.services.export.styles import add_sheet, sheet_title

#: `result_data` keys in the order a reader wants them, across every activity
#: type. Anything an activity wrote that is not listed still gets a column,
#: appended after these — `Activity.config` is free-form and nothing validates
#: what lands in `result_data`.
_KNOWN_RESULT_KEYS = (
    "completion_time_seconds",
    "achieved_points",
    "max_possible_points",
    "success",
    "assigned_points",
    "reasoning",
    "points",
    "result",
    "opponent_team_id",
    "match_duration_seconds",
    "completed",
)


@dataclass(frozen=True)
class ResultsIndex:
    """Where the flat Results sheet's columns and rows landed.

    Aggregates on other sheets need these to build their ranges; nothing may
    hard-code them, since the counter columns vary per event.
    """

    first_row: int
    last_row: int
    team_col: int
    checkpoint_col: int
    activity_col: int
    score_col: int

    @property
    def is_empty(self) -> bool:
        return self.last_row < self.first_row

    def range(self, column: int) -> str:
        return f.ref(f.RESULTS_SHEET, self.first_row, self.last_row, column)


def _result_data_keys(results: list[ActivityResult]) -> list[str]:
    """Every `result_data` key actually written, known ones first."""
    present: set[str] = set()
    for r in results:
        data = getattr(r, "result_data", None)
        if isinstance(data, dict):
            present.update(str(k) for k in data)
    present.discard("notes")
    ordered = [k for k in _KNOWN_RESULT_KEYS if k in present]
    ordered += sorted(present - set(_KNOWN_RESULT_KEYS))
    return ordered


def _unit(data: Any, key: str, *, bonus: bool = False) -> str:
    """Marks a column that holds priced points rather than occurrences."""
    return "" if data.key_is_counted(key, bonus=bonus) else " (pontos)"


def _result_data(result: ActivityResult, key: str) -> Any:
    data = getattr(result, "result_data", None)
    return data.get(key) if isinstance(data, dict) else None


def build_overview_sheet(wb: Workbook, ctx: EventReportContext) -> None:
    """Identity, window, counts, and which mechanics were switched on."""
    event = ctx.event
    data = ctx.results
    rows: list[list[Any]] = [
        ["Event", getattr(event, "name", "") or "—"],
        ["Event type", getattr(event, "event_type", "") or "—"],
        ["Slug", getattr(event, "slug", "") or "—"],
        ["Start", getattr(event, "start_time", None)],
        ["End", getattr(event, "end_time", None)],
        ["Teams", len(data.teams)],
        ["Checkpoints planned", len(data.checkpoints)],
        ["Checkpoints used", len(data.used_checkpoints())],
        ["Activities", len(ctx.content.activities)],
        ["Results recorded", len(data.results)],
        ["Staff assignments", len(ctx.roster.staff_assignments)],
        ["Guide assignments", len(ctx.roster.guide_assignments)],
        ["Team members", len(ctx.roster.members)],
        ["Score edits logged", len(ctx.audit.evaluations)],
        ["Admin actions logged", len(ctx.audit.audit_log)],
        # A false reading here means a score was changed without the team
        # total being recomputed, which is worth stating rather than hiding.
        ["Totals reconcile", "yes" if ctx.totals_reconcile() else "NO — totals disagree"],
    ]
    rows += [[label, "on" if enabled else "off"] for label, enabled in ctx.enabled_settings()]

    add_sheet(wb, "Overview", ["Field", "Value"], rows, center_from=3)


def build_results_sheet(wb: Workbook, ctx: EventReportContext) -> ResultsIndex:
    """One row per result — the source every other aggregate reads."""
    data = ctx.results
    detail_keys = _result_data_keys(data.results)
    headers, detail_start = _result_headers(data, detail_keys)
    rows = _result_rows(ctx, detail_keys)

    notes_col = detail_start + len(detail_keys)
    reasoning_cols = {detail_start + i for i, k in enumerate(detail_keys) if k == "reasoning"}
    add_sheet(
        wb,
        f.RESULTS_SHEET,
        headers,
        rows,
        center_from=3,
        text_cols={notes_col, *reasoning_cols},
    )

    return ResultsIndex(
        first_row=2,
        last_row=len(rows) + 1,
        team_col=1,
        checkpoint_col=2,
        activity_col=4,
        score_col=6,
    )


def _result_headers(data: Any, detail_keys: list[str]) -> tuple[list[str], int]:
    headers = ["Team", "Checkpoint", "Order", "Activity", "Activity Type", "Final Score"]
    headers += ["Completed", "Judgment", "Completed At", "Extra Shots"]
    headers += [f"P: {data.key_label(key)}{_unit(data, key)}" for key in data.penalty_keys_used]
    headers += [
        f"B: {data.key_label(key)}{_unit(data, key, bonus=True)}" for key in data.bonus_keys_used
    ]
    detail_start = len(headers) + 1
    headers += [f"D: {key}" for key in detail_keys] + ["Notes", "Media", "Edits"]
    return headers, detail_start


def _result_rows(ctx: EventReportContext, detail_keys: list[str]) -> list[list[Any]]:
    data = ctx.results
    checkpoints = {checkpoint.id: checkpoint for checkpoint in data.checkpoints}
    edits: dict[int, int] = {}
    for history in ctx.audit.evaluations:
        edits[history.result_id] = edits.get(history.result_id, 0) + 1

    rows: list[list[Any]] = []
    for result in data.results:
        activity = getattr(result, "activity", None)
        checkpoint = getattr(activity, "checkpoint", None)
        checkpoint_id = getattr(checkpoint, "id", None) if checkpoint else None
        checkpoint = checkpoints.get(checkpoint_id) if checkpoint_id is not None else None
        rows.append([
            data.team_name(result.team_id),
            getattr(checkpoint, "name", "") if checkpoint else "",
            getattr(checkpoint, "order", "") if checkpoint else "",
            getattr(activity, "name", "") if activity else "",
            getattr(activity, "activity_type", "") if activity else "",
            result.final_score,
            "yes" if result.is_completed else "no",
            getattr(result, "judgment_status", None) or "",
            getattr(result, "completed_at", None),
            int(getattr(result, "extra_shots", 0) or 0),
            *[data.key_amount(result, key) for key in data.penalty_keys_used],
            *[data.key_amount(result, key, bonus=True) for key in data.bonus_keys_used],
            *[_result_data(result, key) for key in detail_keys],
            result_notes(result),
            len(getattr(result, "media_urls", None) or []),
            edits.get(getattr(result, "id", -1), 0),
        ])
    return rows


def build_overall_sheet(wb: Workbook, ctx: EventReportContext) -> None:
    """One row per team, one column per checkpoint used, plus reconciliation.

    The checkpoint sum alone is not the team's score: hint and skip costs and
    manual awards are applied against `Team.total` and never against a
    checkpoint. Showing the subtotal, the adjustments and the recorded total
    side by side is what makes a disagreement visible instead of silent.
    """
    data = ctx.results
    checkpoints = data.used_checkpoints()
    show_versus = data.has_versus

    headers = ["Team"]
    if show_versus:
        headers.append("Versus Pair")
    first_cp_col = len(headers) + 1
    headers += [f"{i}. {cp.name}" for i, cp in enumerate(checkpoints, start=1)]
    last_cp_col = len(headers)
    headers += ["Checkpoint Subtotal", "Adjustments", "Recorded Total", "Rank"]

    subtotal_col = last_cp_col + 1
    total_col = last_cp_col + 3

    rows: list[list[Any]] = []
    for offset, team in enumerate(data.teams):
        excel_row = offset + 2
        row: list[Any] = [team.name]
        if show_versus:
            row.append(data.opponent_of.get(team.id, ""))
        for cp in checkpoints:
            row.append(
                data.cp_score.get((team.id, cp.id), 0.0)
                if data.team_attended(team.id, cp.id)
                else ABSENT_XLSX
            )
        row.append(f.row_sum(excel_row, first_cp_col, last_cp_col))
        row.append(ctx.team_adjustments(team.id))
        row.append(ctx.team_recorded_total(team.id))
        row.append(
            f.rank(
                f"{get_column_letter(total_col)}{excel_row}",
                total_col,
                2,
                len(data.teams) + 1,
            )
        )
        rows.append(row)

    last_row = len(rows) + 1
    total_row: list[Any] | None = None
    if rows:
        total_row = ["Total"] + [""] * (first_cp_col - 2)
        total_row += [f.col_sum(c, 2, last_row) for c in range(first_cp_col, subtotal_col + 3)]
        total_row.append("")

    add_sheet(
        wb,
        "Overall",
        headers,
        rows,
        center_from=first_cp_col,
        total_row=total_row,
    )


def build_checkpoint_sheets(wb: Workbook, ctx: EventReportContext) -> None:
    """One sheet per checkpoint that anything actually happened at."""
    data = ctx.results
    columns = checkpoint_columns(data, extended=True)
    taken: set[str] = set(wb.sheetnames)

    for index, checkpoint in enumerate(data.checkpoints, start=1):
        if not data.checkpoint_used(checkpoint.id):
            continue
        rows = [
            checkpoint_row(data, columns, team, checkpoint, absent=ABSENT_XLSX)
            for team in data.teams
        ]
        add_sheet(
            wb,
            sheet_title(checkpoint.name, taken, prefix=f"{index}. "),
            [c.header_en for c in columns],
            rows,
            center_from=2,
            text_cols=_text_columns(columns),
            total_row=_checkpoint_total_row(columns, len(rows)),
        )


def _text_columns(columns: list[ReportColumn]) -> set[int]:
    return {i for i, c in enumerate(columns, start=1) if c.header_en == "Notes"}


def _checkpoint_total_row(columns: list[ReportColumn], body_rows: int) -> list[Any] | None:
    """A totals row summing every numeric column."""
    if not body_rows:
        return None
    last_row = body_rows + 1
    return ["Total"] + [
        f.col_sum(index, 2, last_row) if column.numeric else ""
        for index, column in enumerate(columns[1:], start=2)
    ]
